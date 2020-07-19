import random
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

##################################################################################

OrderInLaw = ["편","장","절","관","조","항","호","목"]
paragraph_symbol = ['①','②','③','④','⑤','⑥','⑦','⑧','⑨','⑩','⑪','⑫','⑬','⑭','⑮','⑯','⑰','⑱','⑲','⑳']
item_symbol = ['가','나','다','라','마','바','사','아','자','차','카','타','파','하']

"""
https://www.lawmaking.go.kr/lmKnlg/jdgStd/info?astSeq=96&astClsCd=
대한민국의 법령은 위와 같이 Index를 나눈다. 
Article은 조, Paragraph는 항, Subparagraph는 호, Item은 목을 영어로 표기한 것이다.  
"""


##################################################################################

class Problem(object):
    
    def __init__(self):

        self.Whole_text = []
        self.Problem_text = []
        self.Last_class_name = ["","","",""]
        self.Last_class = 0
 
    def put_data(self,Line_Type,Text):
        """
        데이터 입력함수. 
        조, 항, 호, 목의 본문은 각각 하나의 Line에 할당된다.
        즉, Line_Type은 "제5조의2", "제1항" 같은 Str 타입의 정보를 담고 있다.
        다만, 하나의 Line에 조,항이 같이 포함되는 등, 복잡한 경우가 있어 주의해야한다.
        """ 
        (val_A, val_P, val_S, val_I) = Line_Type

        # 문제 class에서는 조문을 구분하는 가장 높은 단위는 '조'다. 
        if val_A:
            self.Last_class_name[0]= val_A

        # 제1항 다음에 나오는 본문이 제2항이라면, present class는 1, last class도 1. 
        # 제1호 다목 다음에 나오는 본문이 2호라면, present class는 2, last class는 3. 
        for i in range(len(Line_Type)):
            if Line_Type[i]:
                present_class = i

        #제1항제1호다목 -> 제1항제2호 ;
        #["제10조","제1항","제1호",다목] -> ["제10조","제1항","제2호",""]
        self.Last_class_name[present_class] = Line_Type[present_class]
        if present_class <= self.Last_class:
            for i in range(len(self.Last_class_name)):
                if i > present_class:
                    self.Last_class_name[i] = ""

        total_class_name = ""
        for name in self.Last_class_name:
            total_class_name += str(name)
        self.Whole_text.append([total_class_name,Text])        
        self.Last_class = present_class 
    
    def blank(self,number,mode=0,keyword = None):
        self.Problem_text = All_Text(self.Whole_text)        
        if mode == 0:
            Problem_text_list1 = self.Problem_text.split(" ")
            Problem_text_list2 = [word for word in Problem_text_list1 if word]
            Problem_text_list2 = random_blank(Problem_text_list2,number)
            self.Problem_text = " ".join(Problem_text_list2)
            return self.Problem_text
        elif mode == 1:
            return 0
        else :
            return 0


##################################################################################

def All_Text(The_list):
    All = ""
    for text in The_list:
        All += str(text[1]) + "\n"
    return All

def All_Text2(The_list):
    All = ""
    for text in The_list:
        All += str(text) + "\n"
    return All

#################################################################################

def random_sublist(the_list,number):
    order_list = []
    sublist = []
    max_index = len(the_list)-1
    
    def drawing():
        n = random.randint(0,max_index)
        if n in order_list:
            return drawing()
        else :
            order_list.append(n)
            return order_list
    for i in range(number):
        drawing()
    for n in order_list:
        sublist.append(the_list[n])
    return order_list, sublist

def random_blank(the_list,number):
    order_list, sublist = random_sublist(the_list,number)
    for i in range(number):
        word = sublist[i]
        blank = ""
        for j in range(len(word)):
            blank += " "
        blank = "(" + blank + ")"
        sublist[i] = blank
    for i in range(number):
        the_list[order_list[i]] = sublist[i]
    return the_list

#################################################################################

def replaceRight(original, old, new, count_right):
    """
    https://yuddomack.tistory.com/entry/파이썬-replace-문자열-제거-수정변환
    ...에서 함수의 내용을 참조.
    Str class에서 replace 함수가 좌측부터 한다는 것과 '우측부터 하는 함수'가 없는 문제가 있었다.
    다른 제작자분이 직접 만들어 놓으셨다. 검색해서 표출이 되지 않았다면 직접 만들 심산이었으나
    어깨가 가벼워졌다. 
    """
    repeat = 0
    text = original
    old_len = len(old)

    count_find = original.count(old)
    if count_right > count_find :
        repeat = count_find
    else :
        repeat = count_right

    while(repeat):
        find_index = text.rfind(old)
        text = text[:find_index] + new + text[find_index+old_len:]

        repeat -= 1

    return text 


##################################################################################

def Article_finder(text_list):
    # ['제3조',...]와 같이 구분된 텍스트에서 '제'를 찾고, 제3조의2 같은 신설조항을 구분해준다. 
    text = text_list[0]
    total = ""
    if text[0] == '제':
        for i in range(3):
            if OrderInLaw[i] == text[-1]:
                return 0
        for letter in text:
            try :
                if (letter == "의" or letter == OrderInLaw[4] or int(letter)):
                    total += letter
            except :
                continue
        while total[-1] == "의":
            total = replaceRight(total,'의','',1)
        return "제" + total 
    else :
        return 0 


def Paragraph_finder(text_list):
    # 텍스트에서 처음 표기된 특수문자를 찾고 값을 되돌린다.(예: '①','②'...)
    total = ""
    for unit in text_list:
        if unit in paragraph_symbol:
            total = unit
            break 
        else :
            continue
    if total == "":
        return 0
    else :
        return '제' + unit + '항'
    

def Subparagraph_finder(text_list):
    # '1.' 같이 적힌 호 표기를 찾는다. 온점을 찾은 다음, int 값으로 변환가능한지 본다.  
    total = ""
    text = text_list[0]
    total = text[:-1]
    if text[-1] == "." :
        try :
            int(text[:-1])
        except :
            return 0
        return "제" + total + "호"
    else :
        return 0 

def Item_finder(text_list):
    # '가.' 같이 적힌 목 표기를 찾는다. 온점을 찾은 다음, 가나다라...중 하나인지 확인한다. 
    total = ""
    text = text_list[0]
    total = text[:-1]
    if text[-1] == "." and (total in item_symbol):
        return total + "목"
    else :
        return 0 

##################################################################################
"""
templit은 읽어들일 조문이 들어있는 워드 파일.
Result는 출제된 문제가 들어있는 워드 파일.
NP는 number of problem, 즉 문제의 개수.
Keyword는 출제하고 싶은 문제와 관련된 키워드.
"""


templit = Document('Test1.docx')
Result = Document('Today.docx')
NP = 3
Keyword = None 

##################################################################################

"""
https://m.blog.naver.com/anakt/221842622079
다음 조항으로 넘어갈 때 '문제 유닛'을 클래스 선언하고, 이전 '문제 유닛'을 '문제 리스트'에 넣는다.
"""
def data_acquisition(templit):
    Problem_list = []
    Problem_unit = None
    
    for row, paragraph in enumerate(templit.paragraphs):
        if row > 1000:
            break

        line1 = paragraph.text.split(' ')
        line2 = [word for word in line1 if word]
        if line2 == []:
            continue
        
        Line_Type = (val_A, val_P, val_S, val_I) = (Article_finder(line2), Paragraph_finder(line2), Subparagraph_finder(line2), Item_finder(line2)) 
        if Line_Type == (0,0,0,0):
            continue
        
        if val_A :
            if Problem_unit:
                Problem_list.append(Problem_unit)
            Problem_unit = Problem()
        Problem_unit.put_data(Line_Type,paragraph.text)
        
    return Problem_list

##################################################################################    
"""
여기 함수들은 출제방식을 나타낸다.
Set_Exam_Random()는 리스트에 있는 조문을 랜덤하게 뽑기만 한다. 
"""

def Set_Exam_AllRandom(Result,Problem_list,NP):

    for i in range(NP):
        new_problem = random.choice(Problem_list)
        Problem_list.remove(new_problem)
        new_problem.blank(3,mode=0,keyword=None)
        Result.add_paragraph(new_problem.Problem_text)
    Result.save('Today.docx')

##################################################################################
        
Problem_list = data_acquisition(templit)
Set_Exam_AllRandom(Result,Problem_list,3)




    
    
